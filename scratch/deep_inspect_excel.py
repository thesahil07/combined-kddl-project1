import openpyxl

for fname in ['KDDL_Assessment_Dials_1_Demo_FY_2026_Cycle_1.xlsx', 'KDDL_Assessment_Gaps_FY2026 (8).xlsx']:
    path = r'c:\Users\VVNT\Downloads\\' + fname
    wb = openpyxl.load_workbook(path)
    print(f'==============================')
    print(f'FILE: {fname}')
    print(f'==============================')
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f'\n--- Sheet: {sname} (Rows: {ws.max_row}, Cols: {ws.max_column}) ---')
        # Print non-empty sample rows
        for r in range(1, min(ws.max_row + 1, 15)):
            vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            # filter None at the end
            while vals and vals[-1] is None:
                vals.pop()
            print(f'Row {r}:', vals[:10])
