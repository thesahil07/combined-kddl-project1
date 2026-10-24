import openpyxl

path = r'c:\Users\VVNT\Downloads\KDDL_Assessment_Dials_1_Demo_FY_2026_Cycle_1.xlsx'
wb = openpyxl.load_workbook(path)
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f'=== Sheet: {sname} (Rows: {ws.max_row}) ===')
    for r in range(1, min(ws.max_row + 1, 10)):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        print(f'Row {r}:', row_vals)
