import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\VVNT\Downloads\KDDL_Assessment_Gaps_FY2026 (8).xlsx')
for sname in ['Success Behaviour', 'Functional Skills']:
    ws = wb[sname]
    print(f'*** SHEET: {sname} ***')
    print('Row 3:', [ws.cell(3, c).value for c in range(1, ws.max_column + 1) if ws.cell(3, c).value])
    print('Row 4:', [ws.cell(4, c).value for c in range(1, ws.max_column + 1)])
    print('Row 5:', [ws.cell(5, c).value for c in range(1, ws.max_column + 1)])
