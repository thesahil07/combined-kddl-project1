import io, openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text
import dotenv
dotenv.load_dotenv('backend/.env')
from app.presentation.api.dependencies import initialize_database_engine, get_session

initialize_database_engine()
s = next(get_session())

tenant_id = 1
division_id = None
cycle_id = 1
employee_name = None

# Query divisions
divisions = s.execute(text("SELECT division_id, code, name FROM divisions WHERE tenant_id=:tenant_id"), {"tenant_id": tenant_id}).mappings().all()

# Query cycles
cycles = s.execute(text("""
    SELECT cycle_id, name FROM assessment_cycles
    WHERE tenant_id=:tenant_id AND status IN ('NOT_STARTED', 'IN_PROGRESS')
    ORDER BY period_start DESC
"""), {"tenant_id": tenant_id}).mappings().all()
selected_cycle_id = cycle_id or (cycles[0]["cycle_id"] if cycles else None)

# Division label
division_label = "All Entities"
if division_id is not None:
    div_row = s.execute(
        text("SELECT name FROM divisions WHERE tenant_id=:tenant_id AND division_id=:division_id"),
        {"tenant_id": tenant_id, "division_id": division_id},
    ).scalar_one_or_none()
    if div_row:
        division_label = str(div_row).replace(" - Demo", "")

print(f"Generating report for Division: {division_label}, Cycle: {selected_cycle_id}")

# 1. Fetch all employees in scope
emp_query = text("""
    SELECT e.employee_id, e.employee_code, e.full_name AS employee_name,
           REPLACE(d.name, ' - Demo', '') AS division_name,
           COALESCE(m.full_name, 'Not assigned') AS manager_name,
           COALESCE(f.name, '-') AS function_name,
           COALESCE(dep.name, '-') AS department_name,
           COALESCE(r.title, '-') AS role_title,
           COALESCE(gb.band_code, '-') AS grade,
           COALESCE(e.designation, '-') AS designation,
           CASE WHEN behaviour.form_id IS NULL THEN 'PENDING_WITH_EMPLOYEE'
                WHEN behaviour.behaviour_self_submitted_at IS NULL THEN 'PENDING_WITH_EMPLOYEE'
                ELSE 'COMPLETED' END AS behaviour_self_status,
           CASE WHEN behaviour.form_id IS NULL THEN 'PENDING_WITH_MANAGER'
                WHEN behaviour.behaviour_manager_submitted_at IS NULL THEN 'PENDING_WITH_MANAGER'
                ELSE 'COMPLETED' END AS behaviour_manager_status,
           CASE WHEN functional.form_id IS NULL AND EXISTS (
                    SELECT 1 FROM functional_skills fs
                    JOIN desired_level_matrix dlm ON dlm.tenant_id=fs.tenant_id
                        AND dlm.division_id=e.division_id AND dlm.item_type='SKILL'
                        AND dlm.item_id=fs.skill_id AND dlm.department_id <=> e.department_id
                        AND dlm.grade_band_id=e.grade_band_id
                    WHERE fs.tenant_id=e.tenant_id AND fs.division_id=e.division_id
                        AND fs.function_id=e.function_id AND fs.department_id <=> e.department_id
                ) THEN 'PENDING_WITH_EMPLOYEE'
                WHEN functional.form_id IS NULL THEN 'NOT_CONFIGURED'
                WHEN functional.functional_self_submitted_at IS NULL THEN 'PENDING_WITH_EMPLOYEE'
                ELSE 'COMPLETED' END AS functional_self_status,
           CASE WHEN functional.form_id IS NULL AND EXISTS (
                    SELECT 1 FROM functional_skills fs
                    JOIN desired_level_matrix dlm ON dlm.tenant_id=fs.tenant_id
                        AND dlm.division_id=e.division_id AND dlm.item_type='SKILL'
                        AND dlm.item_id=fs.skill_id AND dlm.department_id <=> e.department_id
                        AND dlm.grade_band_id=e.grade_band_id
                    WHERE fs.tenant_id=e.tenant_id AND fs.division_id=e.division_id
                        AND fs.function_id=e.function_id AND fs.department_id <=> e.department_id
                ) THEN 'PENDING_WITH_MANAGER'
                WHEN functional.form_id IS NULL THEN 'NOT_CONFIGURED'
                WHEN functional.functional_manager_submitted_at IS NULL THEN 'PENDING_WITH_MANAGER'
                ELSE 'COMPLETED' END AS functional_manager_status,
           COALESCE(g.behaviour_gaps, 0) AS behaviour_gaps,
           COALESCE(g.functional_gaps, 0) AS functional_gaps
    FROM employees e
    JOIN divisions d ON d.division_id=e.division_id AND d.tenant_id=e.tenant_id
    JOIN entities en ON en.entity_id=d.entity_id AND en.tenant_id=d.tenant_id
    LEFT JOIN employees m ON m.employee_id=e.manager_employee_id AND m.tenant_id=e.tenant_id
    LEFT JOIN roles r ON r.role_id=e.role_id AND r.tenant_id=e.tenant_id AND r.division_id=e.division_id
    LEFT JOIN functions f ON f.function_id=e.function_id
    LEFT JOIN departments dep ON dep.department_id=e.department_id
    LEFT JOIN grade_bands gb ON gb.grade_band_id=e.grade_band_id
    LEFT JOIN assessment_forms behaviour ON behaviour.tenant_id=e.tenant_id AND behaviour.employee_id=e.employee_id
        AND behaviour.cycle_id=:cycle_id AND behaviour.assessment_type='BEHAVIOUR'
    LEFT JOIN assessment_forms functional ON functional.tenant_id=e.tenant_id AND functional.employee_id=e.employee_id
        AND functional.cycle_id=:cycle_id AND functional.assessment_type='SKILL'
    LEFT JOIN (
        SELECT af.employee_id,
            SUM(CASE WHEN ali.item_type='BEHAVIOUR' AND ali.gap < 0 THEN 1 ELSE 0 END) AS behaviour_gaps,
            SUM(CASE WHEN ali.item_type='SKILL' AND ali.gap < 0 THEN 1 ELSE 0 END) AS functional_gaps
        FROM assessment_forms af JOIN assessment_line_items ali ON ali.form_id=af.form_id
        WHERE af.tenant_id=:tenant_id AND af.cycle_id=:cycle_id GROUP BY af.employee_id
    ) g ON g.employee_id=e.employee_id
    WHERE e.tenant_id=:tenant_id AND e.status='ACTIVE' AND en.code <> 'CORPORATE'
      AND d.code <> 'TARATEC-LDR'
      AND (:division_id IS NULL OR e.division_id=:division_id)
      AND (:employee_name IS NULL OR LOWER(e.full_name) LIKE :employee_name_like OR LOWER(e.employee_code) LIKE :employee_name_like)
    ORDER BY d.name, e.full_name
""")

employees = s.execute(emp_query, {
    "tenant_id": tenant_id,
    "division_id": division_id,
    "cycle_id": selected_cycle_id,
    "employee_name": employee_name,
    "employee_name_like": f"%{employee_name.lower()}%" if employee_name else "%",
}).mappings().all()

print(f"Total employees: {len(employees)}")

# 2. Fetch all line items for the cycle
line_query = text("""
    SELECT
        af.employee_id,
        ali.line_item_id,
        ali.item_type,
        ali.item_id,
        COALESCE(sb.name, fs.name) AS item_name,
        ali.desired_level,
        ali.self_assessment_level,
        ali.manager_assessment_level,
        ali.gap
    FROM assessment_forms af
    JOIN assessment_line_items ali ON ali.form_id = af.form_id
    LEFT JOIN success_behaviours sb ON ali.item_type = 'BEHAVIOUR' AND sb.behaviour_id = ali.item_id
    LEFT JOIN functional_skills fs ON ali.item_type = 'SKILL' AND fs.skill_id = ali.item_id
    WHERE af.tenant_id = :tenant_id AND af.cycle_id = :cycle_id
    ORDER BY af.employee_id, ali.item_type, ali.line_item_id
""")

line_rows = s.execute(line_query, {
    "tenant_id": tenant_id,
    "cycle_id": selected_cycle_id,
}).mappings().all()

print(f"Total line items: {len(line_rows)}")

# Group line items by employee_id and item_type
emp_behaviours = {}  # emp_id -> {sb_id: row}
emp_skills = {}      # emp_id -> list of skill rows

for r in line_rows:
    emp_id = r["employee_id"]
    if r["item_type"] == "BEHAVIOUR":
        if emp_id not in emp_behaviours:
            emp_behaviours[emp_id] = {}
        emp_behaviours[emp_id][r["item_id"]] = r
    elif r["item_type"] == "SKILL":
        if emp_id not in emp_skills:
            emp_skills[emp_id] = []
        emp_skills[emp_id].append(r)

print(f"Employees with behaviour items: {len(emp_behaviours)}")
print(f"Employees with skill items: {len(emp_skills)}")

# The 5 universal success behaviours in order (IDs 6 to 10)
behaviours_meta = [
    (6, "Behaviour 1: Ownership & Execution Discipline"),
    (7, "Behaviour 2: Problem Solving & Critical Thinking"),
    (8, "Behaviour 3: Collaboration & Influencing"),
    (9, "Behaviour 4: Detail Orientation & Quality Focus"),
    (10, "Behaviour 5: Learning Agility & Coaching Mindset"),
]

wb = openpyxl.Workbook()
ws_sb = wb.active
ws_sb.title = "Success Behaviours"
ws_fs = wb.create_sheet("Functional Skills")
ws_gaps = wb.create_sheet("All Assessment Gaps")
ws_status = wb.create_sheet("Assessment Status Summary")

# STYLES
banner_font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
subtitle_font = Font(name="Calibri", size=10, italic=True, color="065F46")
super_header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
col_header_font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
regular_font = Font(name="Calibri", size=9)
bold_font = Font(name="Calibri", size=9, bold=True)

navy_banner_fill = PatternFill(start_color="162238", end_color="162238", fill_type="solid")
teal_banner_fill = PatternFill(start_color="0F4C5C", end_color="0F4C5C", fill_type="solid")
subtitle_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")

navy_super_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
blue_super_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
teal_super_fill = PatternFill(start_color="115E59", end_color="115E59", fill_type="solid")

col_header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
teal_col_header_fill = PatternFill(start_color="134E4A", end_color="134E4A", fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
red_font = Font(name="Calibri", size=9, bold=True, color="991B1B")
green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
green_font = Font(name="Calibri", size=9, color="166534")
blue_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
blue_font = Font(name="Calibri", size=9, color="1E40AF")
amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
amber_font = Font(name="Calibri", size=9, color="92400E")

# -------------------------------------------------------------
# SHEET 1: Success Behaviours (Horizontal Matrix)
# Total columns: 9 + 5 * 4 = 29 (A to AC)
# -------------------------------------------------------------
ws_sb.views.sheetView[0].showGridLines = True
ws_sb.freeze_panes = "J5"

# Row 1: Title Banner
ws_sb.merge_cells(start_row=1, start_column=1, end_row=1, end_column=29)
title_cell = ws_sb.cell(row=1, column=1, value=f"ASK Gap Assessment - {division_label} - Success Behaviours")
title_cell.font = banner_font
title_cell.fill = navy_banner_fill
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws_sb.row_dimensions[1].height = 32

for c in range(2, 30):
    ws_sb.cell(row=1, column=c).fill = navy_banner_fill

# Row 2: Subtitle
ws_sb.merge_cells(start_row=2, start_column=1, end_row=2, end_column=29)
sub_cell = ws_sb.cell(row=2, column=1, value="All five Success Behaviours are universal across roles. Gap = Manager Assessment Level minus Desired Level.")
sub_cell.font = subtitle_font
sub_cell.fill = subtitle_fill
sub_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws_sb.row_dimensions[2].height = 22

for c in range(2, 30):
    ws_sb.cell(row=2, column=c).fill = subtitle_fill

# Row 3: Super-headers
# Columns 1 to 9: Employee Information
ws_sb.merge_cells(start_row=3, start_column=1, end_row=3, end_column=9)
emp_info_cell = ws_sb.cell(row=3, column=1, value="Employee Information")
emp_info_cell.font = super_header_font
emp_info_cell.fill = navy_super_fill
emp_info_cell.alignment = Alignment(horizontal="center", vertical="center")
for c in range(1, 10):
    ws_sb.cell(row=3, column=c).fill = navy_super_fill
    ws_sb.cell(row=3, column=c).border = thin_border

# Columns 10 to 29: 5 Behaviours (4 cols each)
for b_idx, (b_id, b_name) in enumerate(behaviours_meta):
    start_c = 10 + b_idx * 4
    end_c = start_c + 3
    ws_sb.merge_cells(start_row=3, start_column=start_c, end_row=3, end_column=end_c)
    b_cell = ws_sb.cell(row=3, column=start_c, value=b_name)
    b_cell.font = super_header_font
    b_cell.fill = blue_super_fill
    b_cell.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(start_c, end_c + 1):
        ws_sb.cell(row=3, column=c).fill = blue_super_fill
        ws_sb.cell(row=3, column=c).border = thin_border

ws_sb.row_dimensions[3].height = 26

# Row 4: Sub-headers
sb_headers = [
    "#", "Emp Code", "Name", "Manager Name", "Function", "Department", "Role", "Grade", "Designation",
    "Desired Level", "Self Assessment", "Mgr Assessment", "Gap",
    "Desired Level", "Self Assessment", "Mgr Assessment", "Gap",
    "Desired Level", "Self Assessment", "Mgr Assessment", "Gap",
    "Desired Level", "Self Assessment", "Mgr Assessment", "Gap",
    "Desired Level", "Self Assessment", "Mgr Assessment", "Gap",
]

for col_idx, h_text in enumerate(sb_headers, 1):
    c = ws_sb.cell(row=4, column=col_idx, value=h_text)
    c.font = col_header_font
    c.fill = col_header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border

ws_sb.row_dimensions[4].height = 28

# Populate Employee Rows
for emp_idx, emp in enumerate(employees, 1):
    row_idx = 4 + emp_idx
    ws_sb.row_dimensions[row_idx].height = 20
    emp_id = emp["employee_id"]

    # Basic Info
    base_vals = [
        emp_idx,
        emp["employee_code"],
        emp["employee_name"],
        emp["manager_name"],
        emp["function_name"],
        emp["department_name"],
        emp["role_title"],
        emp["grade"],
        emp["designation"],
    ]
    for c_i, val in enumerate(base_vals, 1):
        cell = ws_sb.cell(row=row_idx, column=c_i, value=val)
        cell.font = regular_font
        cell.border = thin_border
        if c_i in (1, 2, 8):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # 5 Behaviours
    b_dict = emp_behaviours.get(emp_id, {})
    for b_idx, (b_id, _) in enumerate(behaviours_meta):
        b_data = b_dict.get(b_id)
        start_c = 10 + b_idx * 4

        dl = b_data["desired_level"] if b_data and b_data["desired_level"] is not None else "-"
        sal = b_data["self_assessment_level"] if b_data and b_data["self_assessment_level"] is not None else "-"
        mal = b_data["manager_assessment_level"] if b_data and b_data["manager_assessment_level"] is not None else "-"
        gap = b_data["gap"] if b_data and b_data["gap"] is not None else "-"

        items = [dl, sal, mal, gap]
        for offset, val in enumerate(items):
            c_pos = start_c + offset
            cell = ws_sb.cell(row=row_idx, column=c_pos, value=val)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

            # Style Gap column
            if offset == 3 and isinstance(val, (int, float)):
                if val < 0:
                    cell.fill = red_fill
                    cell.font = red_font
                elif val == 0:
                    cell.fill = green_fill
                    cell.font = green_font
                elif val > 0:
                    cell.fill = blue_fill
                    cell.font = blue_font

# -------------------------------------------------------------
# SHEET 2: Functional Skills (Horizontal Matrix)
# Total columns: 9 + 7 * 5 = 44 (A to AR)
# -------------------------------------------------------------
ws_fs.views.sheetView[0].showGridLines = True
ws_fs.freeze_panes = "J5"

# Row 1: Title Banner
ws_fs.merge_cells(start_row=1, start_column=1, end_row=1, end_column=44)
fs_title = ws_fs.cell(row=1, column=1, value=f"ASK Gap Assessment - {division_label} - Functional Skills")
fs_title.font = banner_font
fs_title.fill = teal_banner_fill
fs_title.alignment = Alignment(horizontal="center", vertical="center")
ws_fs.row_dimensions[1].height = 32
for c in range(2, 45):
    ws_fs.cell(row=1, column=c).fill = teal_banner_fill

# Row 2: Subtitle
ws_fs.merge_cells(start_row=2, start_column=1, end_row=2, end_column=44)
fs_sub = ws_fs.cell(row=2, column=1, value="Functional Skills are role-specific (up to 7 skills per employee). Gap = Manager Assessment Level minus Desired Level.")
fs_sub.font = subtitle_font
fs_sub.fill = subtitle_fill
fs_sub.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws_fs.row_dimensions[2].height = 22
for c in range(2, 45):
    ws_fs.cell(row=2, column=c).fill = subtitle_fill

# Row 3: Super-headers
ws_fs.merge_cells(start_row=3, start_column=1, end_row=3, end_column=9)
fs_emp_info = ws_fs.cell(row=3, column=1, value="Employee Information")
fs_emp_info.font = super_header_font
fs_emp_info.fill = navy_super_fill
fs_emp_info.alignment = Alignment(horizontal="center", vertical="center")
for c in range(1, 10):
    ws_fs.cell(row=3, column=c).fill = navy_super_fill
    ws_fs.cell(row=3, column=c).border = thin_border

for s_idx in range(7):
    start_c = 10 + s_idx * 5
    end_c = start_c + 4
    ws_fs.merge_cells(start_row=3, start_column=start_c, end_row=3, end_column=end_c)
    s_cell = ws_fs.cell(row=3, column=start_c, value=f"Skill {s_idx + 1}")
    s_cell.font = super_header_font
    s_cell.fill = teal_super_fill
    s_cell.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(start_c, end_c + 1):
        ws_fs.cell(row=3, column=c).fill = teal_super_fill
        ws_fs.cell(row=3, column=c).border = thin_border

ws_fs.row_dimensions[3].height = 26

# Row 4: Sub-headers
fs_headers = ["#", "Emp Code", "Name", "Manager Name", "Function", "Department", "Role", "Grade", "Designation"]
for _ in range(7):
    fs_headers.extend(["Skill Name", "Desired Level", "Self Assessment", "Mgr Assessment", "Gap"])

for col_idx, h_text in enumerate(fs_headers, 1):
    c = ws_fs.cell(row=4, column=col_idx, value=h_text)
    c.font = col_header_font
    c.fill = teal_col_header_fill
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border

ws_fs.row_dimensions[4].height = 28

# Populate Employee Rows for Functional Skills
for emp_idx, emp in enumerate(employees, 1):
    row_idx = 4 + emp_idx
    ws_fs.row_dimensions[row_idx].height = 20
    emp_id = emp["employee_id"]

    base_vals = [
        emp_idx,
        emp["employee_code"],
        emp["employee_name"],
        emp["manager_name"],
        emp["function_name"],
        emp["department_name"],
        emp["role_title"],
        emp["grade"],
        emp["designation"],
    ]
    for c_i, val in enumerate(base_vals, 1):
        cell = ws_fs.cell(row=row_idx, column=c_i, value=val)
        cell.font = regular_font
        cell.border = thin_border
        if c_i in (1, 2, 8):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

    skills_list = emp_skills.get(emp_id, [])
    for s_idx in range(7):
        start_c = 10 + s_idx * 5
        if s_idx < len(skills_list):
            s_data = skills_list[s_idx]
            s_name = s_data["item_name"] or "-"
            dl = s_data["desired_level"] if s_data["desired_level"] is not None else "-"
            sal = s_data["self_assessment_level"] if s_data["self_assessment_level"] is not None else "-"
            mal = s_data["manager_assessment_level"] if s_data["manager_assessment_level"] is not None else "-"
            gap = s_data["gap"] if s_data["gap"] is not None else "-"
        else:
            s_name, dl, sal, mal, gap = "-", "-", "-", "-", "-"

        items = [s_name, dl, sal, mal, gap]
        for offset, val in enumerate(items):
            c_pos = start_c + offset
            cell = ws_fs.cell(row=row_idx, column=c_pos, value=val)
            cell.font = regular_font
            cell.border = thin_border
            if offset == 0:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Style Gap column
            if offset == 4 and isinstance(val, (int, float)):
                if val < 0:
                    cell.fill = red_fill
                    cell.font = red_font
                elif val == 0:
                    cell.fill = green_fill
                    cell.font = green_font
                elif val > 0:
                    cell.fill = blue_fill
                    cell.font = blue_font

# -------------------------------------------------------------
# SHEET 3: All Assessment Gaps (115 rows)
# -------------------------------------------------------------
gaps_headers = [
    "Employee Name", "Employee Code", "Division", "Function", "Department",
    "Role", "Grade", "Manager Name", "Category", "Skill / Behaviour Name",
    "Desired Level (DL)", "Self Assessment Level (SAL)", "Manager Assessment Level (MAL)",
    "Gap (MAL - DL)", "Gap Status",
]
ws_gaps.views.sheetView[0].showGridLines = True
ws_gaps.append(gaps_headers)
ws_gaps.row_dimensions[1].height = 28
ws_gaps.freeze_panes = "A2"
for col_idx in range(1, len(gaps_headers) + 1):
    c = ws_gaps.cell(row=1, column=col_idx)
    c.fill = navy_super_fill
    c.font = col_header_font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border

# Populate gaps from line_rows
emp_map = {e["employee_id"]: e for e in employees}
gap_count = 0
for r in line_rows:
    if r["gap"] is not None and r["gap"] < 0:
        gap_count += 1
        emp = emp_map.get(r["employee_id"], {})
        cat = "Success Behaviour" if r["item_type"] == "BEHAVIOUR" else "Functional Skill"
        row_data = [
            emp.get("employee_name", "-"),
            emp.get("employee_code", "-"),
            emp.get("division_name", "-"),
            emp.get("function_name", "-"),
            emp.get("department_name", "-"),
            emp.get("role_title", "-"),
            emp.get("grade", "-"),
            emp.get("manager_name", "-"),
            cat,
            r["item_name"] or "-",
            r["desired_level"],
            r["self_assessment_level"] if r["self_assessment_level"] is not None else "-",
            r["manager_assessment_level"] if r["manager_assessment_level"] is not None else "-",
            int(r["gap"]),
            f"Negative Gap ({r['gap']})",
        ]
        ws_gaps.append(row_data)
        r_idx = ws_gaps.max_row
        ws_gaps.row_dimensions[r_idx].height = 20
        for col_idx in range(1, len(row_data) + 1):
            cell = ws_gaps.cell(row=r_idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border
            if col_idx in (2, 7, 9, 11, 12, 13, 14, 15):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if col_idx in (14, 15):
                cell.fill = red_fill
                cell.font = red_font

print(f"Populated ws_gaps with {gap_count} gaps")

# -------------------------------------------------------------
# SHEET 4: Assessment Status Summary (148 rows)
# -------------------------------------------------------------
status_headers = [
    "Employee Name", "Employee Code", "Division", "Function", "Department",
    "Role", "Grade", "Manager Name", "Behaviour Self Status", "Behaviour Manager Status",
    "Functional Self Status", "Functional Manager Status", "Behaviour Gaps", "Functional Gaps", "Total Gaps",
]
ws_status.views.sheetView[0].showGridLines = True
ws_status.append(status_headers)
ws_status.row_dimensions[1].height = 28
ws_status.freeze_panes = "A2"
for col_idx in range(1, len(status_headers) + 1):
    c = ws_status.cell(row=1, column=col_idx)
    c.fill = navy_super_fill
    c.font = col_header_font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = thin_border

def format_status(val: str) -> str:
    return str(val).replace("_", " ").title()

for emp in employees:
    tot_gaps = int(emp["behaviour_gaps"]) + int(emp["functional_gaps"])
    row_data = [
        str(emp["employee_name"]),
        str(emp["employee_code"]),
        str(emp["division_name"]),
        str(emp["function_name"]),
        str(emp["department_name"]),
        str(emp["role_title"]),
        str(emp["grade"]),
        str(emp["manager_name"]),
        format_status(emp["behaviour_self_status"]),
        format_status(emp["behaviour_manager_status"]),
        format_status(emp["functional_self_status"]),
        format_status(emp["functional_manager_status"]),
        int(emp["behaviour_gaps"]),
        int(emp["functional_gaps"]),
        tot_gaps,
    ]
    ws_status.append(row_data)
    r_idx = ws_status.max_row
    ws_status.row_dimensions[r_idx].height = 20
    for col_idx in range(1, len(row_data) + 1):
        c_cell = ws_status.cell(row=r_idx, column=col_idx)
        c_cell.font = regular_font
        c_cell.border = thin_border
        if col_idx in (2, 7, 9, 10, 11, 12, 13, 14, 15):
            c_cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c_cell.alignment = Alignment(horizontal="left", vertical="center")

        val_str = str(c_cell.value)
        if val_str == "Completed":
            c_cell.fill = green_fill
            c_cell.font = green_font
        elif "Pending" in val_str:
            c_cell.fill = amber_fill
            c_cell.font = amber_font

        if col_idx in (13, 14, 15) and isinstance(c_cell.value, int) and c_cell.value > 0:
            c_cell.fill = red_fill
            c_cell.font = red_font

# Set Column Widths for all sheets
for ws in [ws_sb, ws_fs, ws_gaps, ws_status]:
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        # Avoid row 1 and 2 for merged banner width calculations
        max_len = 0
        for cell in col:
            if cell.row in (1, 2):
                continue
            v = str(cell.value or "")
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

# Set specific nice widths
ws_sb.column_dimensions["A"].width = 5
ws_sb.column_dimensions["B"].width = 14
ws_sb.column_dimensions["C"].width = 22
ws_sb.column_dimensions["D"].width = 20
ws_sb.column_dimensions["E"].width = 16
ws_sb.column_dimensions["F"].width = 16
ws_sb.column_dimensions["G"].width = 18
ws_sb.column_dimensions["H"].width = 8
ws_sb.column_dimensions["I"].width = 16

ws_fs.column_dimensions["A"].width = 5
ws_fs.column_dimensions["B"].width = 14
ws_fs.column_dimensions["C"].width = 22
ws_fs.column_dimensions["D"].width = 20
ws_fs.column_dimensions["E"].width = 16
ws_fs.column_dimensions["F"].width = 16
ws_fs.column_dimensions["G"].width = 18
ws_fs.column_dimensions["H"].width = 8
ws_fs.column_dimensions["I"].width = 16
for s_i in range(7):
    # Skill Name col is 10 + s_i * 5
    s_col_letter = get_column_letter(10 + s_i * 5)
    ws_fs.column_dimensions[s_col_letter].width = 32

wb.active = ws_sb

buf = io.BytesIO()
wb.save(buf)
print(f"Matrix export generated! Size: {len(buf.getvalue())} bytes")
s.close()
