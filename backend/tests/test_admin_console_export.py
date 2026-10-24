import io
import openpyxl
from fastapi.testclient import TestClient

from app.application.use_cases.get_admin_console import GetAdminConsoleUseCase
from app.domain.interfaces.admin_console_repository import AdminConsoleRepository
from app.main import app
from app.presentation.api.dependencies import (
    TenantScope,
    get_admin_console_use_case,
    require_admin_console_access,
)


class FakeAdminConsoleRepo(AdminConsoleRepository):
    def get_assessment_report(self, division_id=None, cycle_id=None, page=1, page_size=20, employee_name=None):
        return None

    def get_submitted_assessment(self, employee_id: int, cycle_id: int | None = None):
        return None

    def export_assessment_gaps_excel(
        self,
        tenant_id: int,
        division_id: int | None = None,
        cycle_id: int | None = None,
        employee_name: str | None = None,
    ) -> tuple[bytes, str]:
        from openpyxl import Workbook

        wb = Workbook()
        ws_sb = wb.active
        ws_sb.title = "Success Behaviours"

        headers = [
            "Employee Name",
            "Employee Code",
            "Division",
            "Function",
            "Department",
            "Role",
            "Grade",
            "Manager Name",
            "Category",
            "Skill / Behaviour Name",
            "Desired Level (DL)",
            "Self Assessment Level (SAL)",
            "Manager Assessment Level (MAL)",
            "Gap (MAL - DL)",
            "Gap Status",
        ]
        ws_sb.append(headers)

        ws_fs = wb.create_sheet(title="Functional Skills")
        ws_fs.append(headers)

        ws_status = wb.create_sheet(title="Assessment Status Summary")
        ws_status.append(headers)

        sample_row = [
            "Aarav Sharma",
            "EMP001",
            "Precision Engineering",
            "Engineering",
            "Design",
            "Senior Engineer",
            "E3",
            "Rajesh Kumar",
            "Success Behaviour",
            "Problem Solving",
            4,
            3,
            2,
            -2,
            "Negative Gap (-2)",
        ]
        ws_sb.append(sample_row)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "KDDL_Assessment_Gaps_Export_2026.xlsx"


def test_admin_console_export_excel_endpoint():
    fake_repo = FakeAdminConsoleRepo()
    fake_use_case = GetAdminConsoleUseCase(fake_repo)

    app.dependency_overrides[get_admin_console_use_case] = lambda: fake_use_case
    app.dependency_overrides[require_admin_console_access] = lambda: TenantScope(
        tenant_id=1, division_scope_id=None, is_platform_admin=True
    )

    client = TestClient(app)
    response = client.get("/api/v1/admin-console/export-excel?cycle_id=1")

    app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert "Assessment_Gaps_Export" in response.headers["content-disposition"]

    # Verify openpyxl can read the returned workbook
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    sheet_names = wb.sheetnames
    assert "All Assessment Gaps" not in sheet_names
    assert "Success Behaviours" in sheet_names
    assert "Functional Skills" in sheet_names
    assert "Assessment Status Summary" in sheet_names

    ws = wb["Success Behaviours"]
    row1 = [cell.value for cell in ws[1]]
    assert "Employee Name" in row1
    assert "Employee Code" in row1
    assert "Function" in row1
    assert "Department" in row1
    assert "Role" in row1
    assert "Grade" in row1
    assert "Manager Name" in row1
    assert "Desired Level (DL)" in row1
    assert "Self Assessment Level (SAL)" in row1
    assert "Manager Assessment Level (MAL)" in row1
    assert "Gap (MAL - DL)" in row1


def test_admin_console_repository_query_filters_active_employees():
    import inspect
    from app.infrastructure.repositories.admin_console_repository_impl import SqlAlchemyAdminConsoleRepository

    src_export = inspect.getsource(SqlAlchemyAdminConsoleRepository.export_assessment_gaps_excel)
    assert "e.status='ACTIVE'" in src_export or "e.status = 'ACTIVE'" in src_export

    src_report = inspect.getsource(SqlAlchemyAdminConsoleRepository.get_assessment_report)
    assert "e.status='ACTIVE'" in src_report or "e.status = 'ACTIVE'" in src_report

