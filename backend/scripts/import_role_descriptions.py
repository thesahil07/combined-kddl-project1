"""Replace role descriptions from the populated client-data workbook."""

from __future__ import annotations

import argparse
import os
import re
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from sqlalchemy import create_engine, text


TENANT_ID = 1
DEFAULT_WORKBOOK = Path(r"C:\Users\Bhawna Narula\VVNT-Founder\OneDrive - VVNT\WIP_Items\Customers - Existing\PGC\FY26\KDDL\KDDL Data\Talent_Client_Data_Template_Populated_FINAL (1).xlsx")


def clean(value: object) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split())


def normalized(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(value).casefold())


def division_code(value: object) -> str:
    return "D1" if clean(value).casefold() in {"dial 1", "dials 1", "d1"} else "D2"


def worksheet_rows(workbook_path: Path) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Role Descriptions"]
        headers = [clean(cell.value) for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        return [dict(zip(headers, row)) for row in worksheet.iter_rows(min_row=2, values_only=True)
                if clean(row[3])]
    finally:
        workbook.close()


def import_descriptions(workbook_path: Path, apply: bool) -> dict[str, int]:
    load_dotenv()
    engine = create_engine(os.environ["DATABASE_URL"], pool_pre_ping=True)
    source = worksheet_rows(workbook_path)
    with engine.begin() as connection:
        divisions = {row["code"]: int(row["division_id"]) for row in connection.execute(text(
            "SELECT division_id, code FROM divisions WHERE tenant_id=:tenant_id"), {"tenant_id": TENANT_ID}).mappings()}
        functions = {(int(row["division_id"]), normalized(row["name"])): int(row["function_id"]) for row in connection.execute(text(
            "SELECT function_id, division_id, name FROM functions WHERE tenant_id=:tenant_id"), {"tenant_id": TENANT_ID}).mappings()}
        roles = {(int(row["division_id"]), int(row["function_id"]), normalized(row["title"])): int(row["role_id"]) for row in connection.execute(text(
            "SELECT role_id, division_id, function_id, title FROM roles WHERE tenant_id=:tenant_id"), {"tenant_id": TENANT_ID}).mappings()}
        employees = {clean(row["employee_code"]): int(row["employee_id"]) for row in connection.execute(text(
            "SELECT employee_id, employee_code FROM employees WHERE tenant_id=:tenant_id"), {"tenant_id": TENANT_ID}).mappings()}

        records: list[dict[str, object]] = []
        for row in source:
            division_id = divisions[division_code(row["Division Code *"])]
            function_id = functions.get((division_id, normalized(row["Function Name *"])))
            role_id = roles.get((division_id, function_id, normalized(row["Role Title *"]))) if function_id else None
            if role_id is None:
                raise ValueError(f"No role match for {clean(row['Division Code *'])}/{clean(row['Function Name *'])}/{clean(row['Role Title *'])}")
            records.append({
                "role_id": role_id, "people_count": row["People Count"],
                "definition": clean(row["Role Definition *"]), "responsibilities": clean(row["Key Responsibilities *"]),
                "kpis": clean(row["KPIs *"]), "specifications": clean(row["Specifications / Experience *"]),
                "education": clean(row["Education & Certifications *"]), "experience": clean(row["Minimum Relevant Experience *"]),
                "knowledge": clean(row["Knowledge Required *"]), "skills": clean(row["Skills Required *"]),
                "attitudes": clean(row["Attitudes & Behaviours *"]),
                "validator": employees.get(clean(row["Validated By (Employee Code)"])), "validated_at": row["Validated On"],
            })
        if apply:
            connection.execute(text("""DELETE FROM role_descriptions
                WHERE role_id IN (SELECT role_id FROM roles WHERE tenant_id=:tenant_id)"""), {"tenant_id": TENANT_ID})
            connection.execute(text("""INSERT INTO role_descriptions
                (role_id,version,people_count,role_definition,key_responsibilities,kpis,specifications_experience,
                 education_certifications,minimum_relevant_experience,knowledge_required,skills_required,
                 attitudes_behaviours,ai_refined,validated_by_employee_id,validated_at)
                VALUES (:role_id,1,:people_count,:definition,:responsibilities,:kpis,:specifications,:education,
                 :experience,:knowledge,:skills,:attitudes,FALSE,:validator,:validated_at)"""), records)
        else:
            connection.rollback()
    return {"source_descriptions": len(source), "matched_descriptions": len(records), "inserted": len(records) if apply else 0}


def main() -> None:
    parser = argparse.ArgumentParser(description="Import role descriptions from client data.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--apply", action="store_true")
    arguments = parser.parse_args()
    print(import_descriptions(arguments.workbook, arguments.apply))


if __name__ == "__main__":
    main()
