"""Import a tenant-scoped functional-skill mapping workbook.

The workbook supplies function, department, role, grade, and skill/desired-level
pairs. Missing master-data records are created only from the workbook values;
the import is idempotent and can be run again safely.
"""

from __future__ import annotations

import argparse
import math
import sys
from collections.abc import Iterable
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine, text

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.config import settings


SKILL_COLUMNS = tuple(
    (
        "Skill Name" if number == 0 else f"Skill Name.{number}",
        "Desired Level" if number == 0 else f"Desired Level.{number}",
    )
    for number in range(7)
)


def clean(value: object) -> str:
    return "" if value is None or (isinstance(value, float) and math.isnan(value)) else str(value).strip()


def grade_code(value: str) -> str:
    """Normalize workbook bands such as E-4 and LM-1 to database codes."""
    return value.replace("-", "").upper()


def skill_pairs(row: dict[str, object]) -> Iterable[tuple[str, int]]:
    for skill_column, level_column in SKILL_COLUMNS:
        skill = clean(row.get(skill_column))
        level = clean(row.get(level_column))
        if not skill and not level:
            continue
        try:
            normalized_level = int(float(level))
        except ValueError:
            normalized_level = 0
        if not skill or normalized_level not in {1, 2, 3, 4}:
            raise ValueError(f"Invalid skill pair for {clean(row['Function'])}/{clean(row['Role'])}: {skill!r}, {level!r}")
        yield skill, normalized_level


def workbook_rows(workbook_path: Path) -> list[dict[str, object]]:
    """Read the mapping sheet without requiring pandas at runtime."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook.active
    try:
        header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        header_counts: dict[str, int] = {}
        headers: list[str] = []
        for value in header_values:
            header = clean(value)
            occurrence = header_counts.get(header, 0)
            header_counts[header] = occurrence + 1
            headers.append(header if occurrence == 0 else f"{header}.{occurrence}")

        return [
            dict(zip(headers, values, strict=True))
            for values in worksheet.iter_rows(min_row=2, values_only=True)
            if any(clean(value) for value in values)
        ]
    finally:
        workbook.close()


def get_or_create_id(connection, select_sql: str, insert_sql: str, params: dict[str, object]) -> int:
    record_id = connection.execute(text(select_sql), params).scalar_one_or_none()
    if record_id is not None:
        return int(record_id)
    result = connection.execute(text(insert_sql), params)
    return int(result.lastrowid)


def import_mapping(workbook_path: Path, tenant_id: int, division_id: int, dry_run: bool) -> dict[str, int]:
    if not settings.database_url:
        raise RuntimeError("DATABASE_URL is not configured.")

    rows = workbook_rows(workbook_path)
    required = {"Function", "Department", "Role", "Grade"}
    missing = required.difference(rows[0] if rows else set())
    if missing:
        raise ValueError(f"Workbook is missing required columns: {', '.join(sorted(missing))}")

    engine = create_engine(settings.database_url)
    counts = {"rows": 0, "skill_pairs": 0, "skills_created": 0, "desired_levels_created": 0}
    seen_pairs: set[tuple[str, str, str, str, str, int]] = set()

    with engine.begin() as connection:
        tenant_exists = connection.execute(text("SELECT tenant_id FROM tenants WHERE tenant_id = :tenant_id"), {"tenant_id": tenant_id}).scalar_one_or_none()
        if tenant_exists is None:
            raise ValueError(f"Tenant {tenant_id} does not exist.")
        division_exists = connection.execute(text("SELECT division_id FROM divisions WHERE division_id = :division_id AND tenant_id = :tenant_id"), {"division_id": division_id, "tenant_id": tenant_id}).scalar_one_or_none()
        if division_exists is None:
            raise ValueError(f"Division {division_id} does not belong to tenant {tenant_id}.")

        next_code = int(connection.execute(text("SELECT COALESCE(MAX(CAST(SUBSTRING(code, 3) AS UNSIGNED)), 0) FROM functional_skills WHERE tenant_id = :tenant_id AND division_id = :division_id AND code REGEXP '^FS[0-9]+$'"), {"tenant_id": tenant_id, "division_id": division_id}).scalar_one())

        for row in rows:
            function_name, department_name, role_title, band_code = (clean(row[column]) for column in ("Function", "Department", "Role", "Grade"))
            if not all((function_name, department_name, role_title, band_code)):
                raise ValueError("Every mapping row must contain Function, Department, Role, and Grade.")
            counts["rows"] += 1
            params = {"tenant_id": tenant_id, "division_id": division_id, "name": function_name}
            function_id = get_or_create_id(connection, "SELECT function_id FROM functions WHERE tenant_id = :tenant_id AND division_id = :division_id AND name = :name", "INSERT INTO functions (tenant_id, division_id, name) VALUES (:tenant_id, :division_id, :name)", params)
            department_params = {**params, "function_id": function_id, "name": department_name}
            department_id = get_or_create_id(connection, "SELECT department_id FROM departments WHERE tenant_id = :tenant_id AND division_id = :division_id AND function_id = :function_id AND name = :name", "INSERT INTO departments (tenant_id, division_id, function_id, name) VALUES (:tenant_id, :division_id, :function_id, :name)", department_params)
            role_params = {"tenant_id": tenant_id, "function_id": function_id, "title": role_title, "division_id": division_id}
            role_id = get_or_create_id(connection, "SELECT role_id FROM roles WHERE tenant_id = :tenant_id AND division_id = :division_id AND function_id = :function_id AND title = :title ORDER BY role_id LIMIT 1", "INSERT INTO roles (tenant_id, division_id, function_id, title) VALUES (:tenant_id, :division_id, :function_id, :title)", role_params)
            grade_band_id = connection.execute(text("SELECT grade_band_id FROM grade_bands WHERE tenant_id = :tenant_id AND band_code = :band_code"), {"tenant_id": tenant_id, "band_code": grade_code(band_code)}).scalar_one_or_none()
            if grade_band_id is None:
                raise ValueError(f"Grade band {band_code!r} is not configured for tenant {tenant_id}.")

            for skill_name, desired_level in skill_pairs(row):
                pair_key = (function_name, department_name, role_title, band_code, skill_name, desired_level)
                if pair_key in seen_pairs:
                    continue
                seen_pairs.add(pair_key)
                counts["skill_pairs"] += 1
                skill_params = {"tenant_id": tenant_id, "division_id": division_id, "function_id": function_id, "department_id": department_id, "name": skill_name}
                skill_id = connection.execute(text("SELECT skill_id FROM functional_skills WHERE tenant_id = :tenant_id AND division_id = :division_id AND function_id = :function_id AND department_id = :department_id AND name = :name"), skill_params).scalar_one_or_none()
                if skill_id is None:
                    next_code += 1
                    skill_params.update({"code": f"FS{next_code:03d}", "definition": f"Functional capability: {skill_name}."})
                    skill_id = connection.execute(text("INSERT INTO functional_skills (tenant_id, division_id, function_id, department_id, code, name, definition, source) VALUES (:tenant_id, :division_id, :function_id, :department_id, :code, :name, :definition, 'MANUAL')"), skill_params).lastrowid
                    counts["skills_created"] += 1
                level_params = {"tenant_id": tenant_id, "division_id": division_id, "item_id": int(skill_id), "grade_band_id": int(grade_band_id), "role_id": role_id, "department_id": department_id, "desired_level": desired_level}
                existing_level = connection.execute(text("SELECT matrix_id FROM desired_level_matrix WHERE tenant_id = :tenant_id AND division_id = :division_id AND item_type = 'SKILL' AND item_id = :item_id AND grade_band_id = :grade_band_id AND role_id = :role_id"), level_params).scalar_one_or_none()
                if existing_level is None:
                    connection.execute(text("INSERT INTO desired_level_matrix (tenant_id, division_id, item_type, item_id, grade_band_id, role_id, department_id, desired_level) VALUES (:tenant_id, :division_id, 'SKILL', :item_id, :grade_band_id, :role_id, :department_id, :desired_level)"), level_params)
                    counts["desired_levels_created"] += 1
                else:
                    connection.execute(text("UPDATE desired_level_matrix SET department_id = :department_id, desired_level = :desired_level WHERE matrix_id = :matrix_id"), {**level_params, "matrix_id": existing_level})

        if dry_run:
            connection.rollback()
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description="Import functional skill mappings from an Excel workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--tenant-id", type=int, default=1)
    parser.add_argument("--division-id", type=int, default=1)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    print(import_mapping(args.workbook, args.tenant_id, args.division_id, args.dry_run))


if __name__ == "__main__":
    main()
