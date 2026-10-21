"""Import skill-specific L1-L4 definitions from the KDDL KS output workbook."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine, text

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.config import settings

TENANT_ID = 1
CODE = re.compile(r"\b(D0*\d+)\s*-")


def norm_code(value: str) -> str:
    match = CODE.search(value)
    if not match:
        raise ValueError(value)
    return f"D{int(match.group(1)[1:]) :03d}"


def parse(path: Path) -> dict[str, dict[int, str]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    # `ReadOnlyWorksheet.cell()` rescans the source XML each time.  Materialise
    # this modest workbook once so the import remains practical on Windows.
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]

    def value(row: int, column: int) -> str:
        if row < 1 or column < 1 or row > len(rows) or column > len(rows[row - 1]):
            return ""
        return str(rows[row - 1][column - 1] or "")

    result: dict[str, dict[int, str]] = {}
    max_row = len(rows)
    max_column = max((len(row) for row in rows), default=0)
    for row in range(1, max_row + 1):
        for column in range(1, max_column + 1):
            cell_value = value(row, column)
            if not CODE.search(cell_value):
                continue
            code = norm_code(cell_value)
            header_row = next((candidate for candidate in range(row + 1, min(row + 8, max_row) + 1)
                               if "LEVEL 1" in value(candidate, column + 1).upper()), None)
            if header_row is None:
                continue
            levels: dict[int, str] = {}
            for level in range(1, 5):
                parts = []
                for candidate in range(header_row + 1, min(header_row + 4, max_row) + 1):
                    item = value(candidate, column + level).strip()
                    if item:
                        parts.append(item)
                if parts:
                    levels[level] = "\n".join(parts)
            if len(levels) == 4:
                result[code] = levels
    return result


def run(path: Path, apply: bool) -> dict[str, int]:
    entries = parse(path)
    engine = create_engine(settings.database_url)
    counts = {"source_skills": len(entries), "skills_matched": 0, "level_definitions_loaded": 0}
    with engine.connect() as connection:
        transaction = connection.begin()
        try:
            connection.execute(text("""CREATE TABLE IF NOT EXISTS functional_skill_level_definitions (
                definition_id BIGINT UNSIGNED AUTO_INCREMENT PRIMARY KEY,
                tenant_id BIGINT UNSIGNED NOT NULL,
                skill_id BIGINT UNSIGNED NOT NULL,
                level TINYINT NOT NULL,
                description TEXT NOT NULL,
                UNIQUE KEY uq_functional_skill_level (tenant_id, skill_id, level),
                CONSTRAINT fk_functional_skill_level_skill FOREIGN KEY (skill_id) REFERENCES functional_skills(skill_id)
            ) ENGINE=InnoDB"""))
            for code, levels in entries.items():
                skills = connection.execute(text("SELECT skill_id FROM functional_skills WHERE tenant_id=:tenant_id AND code=:code"), {"tenant_id": TENANT_ID, "code": code}).scalars().all()
                if skills:
                    counts["skills_matched"] += len(skills)
                for skill_id in skills:
                    for level, description in levels.items():
                        connection.execute(text("""INSERT INTO functional_skill_level_definitions (tenant_id, skill_id, level, description)
                            VALUES (:tenant_id,:skill_id,:level,:description)
                            ON DUPLICATE KEY UPDATE description=VALUES(description)"""), {"tenant_id": TENANT_ID, "skill_id": skill_id, "level": level, "description": description})
                        counts["level_definitions_loaded"] += 1
            if apply: transaction.commit()
            else: transaction.rollback()
        except Exception:
            transaction.rollback(); raise
    return counts


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    print(run(args.workbook, args.apply))
