"""Link Dials 2 employees to validated role descriptions in the source workbook.

The Dials 2 role workbook identifies the employees covered by each detailed
role-description column.  This script uses those explicit assignments, rather
than guessing from similarly named roles, to repair employee.role_id links.
"""

from __future__ import annotations

import argparse
import os
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from sqlalchemy import create_engine, text


TENANT_ID = 1
DIVISION_ID = 2
VALIDATED_AT = datetime(2026, 4, 1)
DEFAULT_WORKBOOK = Path(
    r"C:\Users\Bhawna Narula\VVNT-Founder\OneDrive - VVNT\WIP_Items\Customers - Existing\PGC\FY26\KDDL\KDDL Data\Copy of 4.2 Role Description Exercise (Dials 2) (Final) -12th Jan 2026.xlsx"
)


def normalise(value: object) -> str:
    return re.sub(r"[^a-z]", "", str(value or "").lower())


def text_value(value: object) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ").strip()


def definition_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", text_value(value).lower())


def employee_matches_people_cell(employee_name: str, people_cell: str) -> bool:
    employee = normalise(employee_name)
    people = normalise(people_cell)
    if employee and employee in people:
        return True
    words = re.findall(r"[a-z]{3,}", employee_name.lower())
    return len(words) >= 2 and words[0] in people and words[-1] in people


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--apply", action="store_true")
    arguments = parser.parse_args()

    load_dotenv()
    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL is not configured")

    workbook = openpyxl.load_workbook(arguments.workbook, data_only=True)
    sheet = workbook["D-2 Roles "]
    engine = create_engine(database_url, pool_pre_ping=True)

    with engine.begin() as connection:
        employees = connection.execute(text("""
            SELECT employee_id, employee_code, full_name, role_id
            FROM employees
            WHERE tenant_id = :tenant_id AND division_id = :division_id AND status = 'ACTIVE'
        """), {"tenant_id": TENANT_ID, "division_id": DIVISION_ID}).mappings().all()
        description_rows = connection.execute(text("""
            SELECT rd.role_id, rd.role_definition
            FROM role_descriptions rd
            JOIN roles r ON r.role_id = rd.role_id
            WHERE r.tenant_id = :tenant_id AND r.division_id = :division_id
              AND rd.validated_at = :validated_at
        """), {
            "tenant_id": TENANT_ID,
            "division_id": DIVISION_ID,
            "validated_at": VALIDATED_AT,
        }).mappings().all()

        descriptions_by_definition: dict[str, list[int]] = defaultdict(list)
        for row in description_rows:
            descriptions_by_definition[definition_key(row["role_definition"])].append(int(row["role_id"]))

        assignment_candidates: dict[int, set[int]] = defaultdict(set)
        for column in range(2, sheet.max_column + 1):
            people_cell = text_value(sheet.cell(3, column).value)
            definition = text_value(sheet.cell(8, column).value)
            role_ids = descriptions_by_definition.get(definition_key(definition), [])
            if not people_cell or not definition or len(role_ids) != 1:
                continue
            for employee in employees:
                if employee_matches_people_cell(str(employee["full_name"]), people_cell):
                    assignment_candidates[int(employee["employee_id"])].add(role_ids[0])

        assignments = {
            employee_id: next(iter(role_ids))
            for employee_id, role_ids in assignment_candidates.items()
            if len(role_ids) == 1
        }
        ambiguous = [
            str(employee["employee_code"])
            for employee in employees
            if len(assignment_candidates.get(int(employee["employee_id"]), set())) > 1
        ]

        print(f"Matched employee-to-description assignments: {len(assignments)} of {len(employees)}")
        if ambiguous:
            print(f"Skipped ambiguous workbook assignments: {', '.join(sorted(ambiguous))}")
        if arguments.apply:
            connection.execute(text("""
                UPDATE employees
                SET role_id = :role_id
                WHERE employee_id = :employee_id AND tenant_id = :tenant_id AND division_id = :division_id
            """), [
                {"employee_id": employee_id, "role_id": role_id,
                 "tenant_id": TENANT_ID, "division_id": DIVISION_ID}
                for employee_id, role_id in assignments.items()
            ])
            print(f"Updated employee role links: {len(assignments)}")


if __name__ == "__main__":
    main()
