"""Incrementally import KDDL master data from a populated client workbook."""

from __future__ import annotations

import argparse
import re
import sys
from datetime import date, datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.config import settings


TENANT_ID = 1
SHEETS = (
    "Functions",
    "Departments",
    "Grade Bands",
    "Roles",
    "Functional Skills",
    "Role Descriptions",
    "Employees",
)


def clean(value: object) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split())


def normalized(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(value).casefold())


def text_value(value: object) -> str:
    """Retain line breaks while turning spreadsheet HTML breaks into text."""
    source = str(value or "").replace("\xa0", " ")
    source = re.sub(r"<br\s*/?>", "\n", source, flags=re.IGNORECASE)
    return "\n".join(" ".join(line.split()) for line in source.splitlines()).strip()


def division_code(value: object) -> str:
    match = re.search(r"([12])", clean(value))
    if match is None:
        raise ValueError(f"Unsupported division value: {value!r}")
    return f"D{match.group(1)}"


def as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else None


def workbook_rows(workbook_path: Path, sheet_name: str) -> list[dict[str, object]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        row_iterator = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(row_iterator)]
        return [
            dict(zip(headers, row, strict=False))
            for row in row_iterator
            if any(clean(value) for value in row)
        ]
    finally:
        workbook.close()


def workbook_sheet_names(workbook_path: Path) -> set[str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        return set(workbook.sheetnames)
    finally:
        workbook.close()


def import_workbook(workbook_path: Path, apply: bool) -> dict[str, int]:
    if not settings.database_url:
        raise RuntimeError("DATABASE_URL is not configured.")
    available_sheets = workbook_sheet_names(workbook_path)
    source = {
        sheet: workbook_rows(workbook_path, sheet) if sheet in available_sheets else []
        for sheet in SHEETS
    }
    counts = {"functions_created": 0, "departments_created": 0, "grade_bands_created": 0,
              "roles_created": 0, "functional_skills_created": 0,
              "functional_skills_updated": 0, "role_descriptions_updated": 0,
              "role_descriptions_created": 0, "employees_updated": 0}
    engine = create_engine(settings.database_url, pool_pre_ping=True)

    with engine.connect() as connection:
        transaction = connection.begin()
        try:
            divisions = {
                str(row["code"]): int(row["division_id"])
                for row in connection.execute(text("""
                    SELECT division_id, code FROM divisions WHERE tenant_id=:tenant_id
                """), {"tenant_id": TENANT_ID}).mappings()
            }
            functions = {
                (int(row["division_id"]), normalized(row["name"])): int(row["function_id"])
                for row in connection.execute(text("""
                    SELECT function_id, division_id, name FROM functions WHERE tenant_id=:tenant_id
                """), {"tenant_id": TENANT_ID}).mappings()
            }
            for row in source["Functions"]:
                division_id = divisions[division_code(row["Division Code *"])]
                function_name = clean(row["Function Name *"])
                lookup_key = (division_id, normalized(function_name))
                if lookup_key not in functions:
                    function_id = connection.execute(text("""
                        INSERT INTO functions (tenant_id, division_id, name)
                        VALUES (:tenant_id, :division_id, :name)
                    """), {"tenant_id": TENANT_ID, "division_id": division_id, "name": function_name}).lastrowid
                    functions[lookup_key] = int(function_id)
                    counts["functions_created"] += 1

            departments = {
                (int(row["division_id"]), int(row["function_id"]), normalized(row["name"])): int(row["department_id"])
                for row in connection.execute(text("""
                    SELECT department_id, division_id, function_id, name FROM departments WHERE tenant_id=:tenant_id
                """), {"tenant_id": TENANT_ID}).mappings()
            }
            for row in source["Departments"]:
                division_id = divisions[division_code(row["Division Code *"])]
                function_id = functions[(division_id, normalized(row["Function Name *"]))]
                department_name = clean(row["Department Name *"])
                lookup_key = (division_id, function_id, normalized(department_name))
                if lookup_key not in departments:
                    department_id = connection.execute(text("""
                        INSERT INTO departments (tenant_id, division_id, function_id, name)
                        VALUES (:tenant_id, :division_id, :function_id, :name)
                    """), {"tenant_id": TENANT_ID, "division_id": division_id,
                              "function_id": function_id, "name": department_name}).lastrowid
                    departments[lookup_key] = int(department_id)
                    counts["departments_created"] += 1

            grade_bands = {
                normalized(row["band_code"]): int(row["grade_band_id"])
                for row in connection.execute(text("""
                    SELECT grade_band_id, band_code FROM grade_bands WHERE tenant_id=:tenant_id
                """), {"tenant_id": TENANT_ID}).mappings()
            }
            for row in source["Grade Bands"]:
                band_code = clean(row["Band Code *"])
                if normalized(band_code) not in grade_bands:
                    grade_band_id = connection.execute(text("""
                        INSERT INTO grade_bands (tenant_id, band_code, band_group, designation_label, sort_order)
                        VALUES (:tenant_id, :band_code, :band_group, :designation_label, :sort_order)
                    """), {"tenant_id": TENANT_ID, "band_code": band_code,
                              "band_group": clean(row["Band Group *"]),
                              "designation_label": clean(row["Designation Label *"]),
                              "sort_order": int(row["Sort Order *"])}).lastrowid
                    grade_bands[normalized(band_code)] = int(grade_band_id)
                    counts["grade_bands_created"] += 1

            roles = {
                (int(row["division_id"]), int(row["function_id"]), normalized(row["title"])): int(row["role_id"])
                for row in connection.execute(text("""
                    SELECT role_id, division_id, function_id, title FROM roles WHERE tenant_id=:tenant_id
                """), {"tenant_id": TENANT_ID}).mappings()
            }
            # A delta may supply a role description for an otherwise omitted role.
            # Create that role first so the description can be imported atomically.
            for row in [*source["Roles"], *source["Role Descriptions"], *source["Employees"]]:
                division_id = divisions[division_code(row["Division Code *"])]
                function_id = functions[(division_id, normalized(row["Function Name *"]))]
                role_title = clean(row.get("Role Title *", row.get("Role Title")))
                lookup_key = (division_id, function_id, normalized(role_title))
                if lookup_key not in roles:
                    role_id = connection.execute(text("""
                        INSERT INTO roles (tenant_id, division_id, function_id, title, description_status)
                        VALUES (:tenant_id, :division_id, :function_id, :title, 'NOT_STARTED')
                    """), {"tenant_id": TENANT_ID, "division_id": division_id,
                              "function_id": function_id, "title": role_title}).lastrowid
                    roles[lookup_key] = int(role_id)
                    counts["roles_created"] += 1

            functional_skills = {
                (int(row["division_id"]), normalized(row["code"])): int(row["skill_id"])
                for row in connection.execute(text("""
                    SELECT skill_id, division_id, code
                    FROM functional_skills
                    WHERE tenant_id=:tenant_id
                """), {"tenant_id": TENANT_ID}).mappings()
            }
            for row in source["Functional Skills"]:
                division_id = divisions[division_code(row["Division Code *"])]
                function_id = functions[(division_id, normalized(row["Function Name *"]))]
                department_name = clean(row.get("Department Name"))
                department_id = (
                    departments[(division_id, function_id, normalized(department_name))]
                    if department_name else None
                )
                code = clean(row["Skill Code *"])
                values = {
                    "tenant_id": TENANT_ID,
                    "division_id": division_id,
                    "function_id": function_id,
                    "department_id": department_id,
                    "code": code,
                    "name": clean(row["Skill Name *"]),
                    "definition": text_value(row["Definition *"]),
                }
                lookup_key = (division_id, normalized(code))
                existing_skill_id = functional_skills.get(lookup_key)
                if existing_skill_id is None:
                    skill_id = connection.execute(text("""
                        INSERT INTO functional_skills
                            (tenant_id, division_id, function_id, department_id, code, name, definition, source)
                        VALUES
                            (:tenant_id, :division_id, :function_id, :department_id, :code, :name, :definition, 'MANUAL')
                    """), values).lastrowid
                    functional_skills[lookup_key] = int(skill_id)
                    counts["functional_skills_created"] += 1
                else:
                    connection.execute(text("""
                        UPDATE functional_skills
                        SET function_id=:function_id, department_id=:department_id,
                            name=:name, definition=:definition, source='MANUAL'
                        WHERE tenant_id=:tenant_id AND skill_id=:skill_id
                    """), {**values, "skill_id": existing_skill_id})
                    counts["functional_skills_updated"] += 1

            employees = {
                normalized(row["employee_code"]): int(row["employee_id"])
                for row in connection.execute(text("""
                    SELECT employee_id, employee_code FROM employees WHERE tenant_id=:tenant_id
                """), {"tenant_id": TENANT_ID}).mappings()
            }
            for row in source["Role Descriptions"]:
                division_id = divisions[division_code(row["Division Code *"])]
                function_id = functions[(division_id, normalized(row["Function Name *"]))]
                role_id = roles[(division_id, function_id, normalized(row["Role Title *"]))]
                values = {
                    "role_id": role_id,
                    "people_count": int(row["People Count"]) if row.get("People Count") is not None else None,
                    "role_definition": text_value(row["Role Definition *"]),
                    "key_responsibilities": text_value(row["Key Responsibilities *"]),
                    "kpis": text_value(row["KPIs *"]),
                    "specifications_experience": text_value(row["Specifications / Experience *"]),
                    "education_certifications": text_value(row["Education & Certifications *"]),
                    "minimum_relevant_experience": text_value(row["Minimum Relevant Experience *"]),
                    "knowledge_required": text_value(row["Knowledge Required *"]),
                    "skills_required": text_value(row["Skills Required *"]),
                    "attitudes_behaviours": text_value(row["Attitudes & Behaviours *"]),
                    "validator": employees.get(normalized(row.get("Validated By (Employee Code)"))),
                    "validated_at": as_date(row.get("Validated On")),
                }
                existing = connection.execute(text("""
                    SELECT role_description_id FROM role_descriptions
                    WHERE role_id=:role_id ORDER BY version DESC, role_description_id DESC LIMIT 1
                """), values).scalar_one_or_none()
                if existing is None:
                    version = connection.execute(text("""
                        SELECT COALESCE(MAX(version), 0) + 1 FROM role_descriptions WHERE role_id=:role_id
                    """), values).scalar_one()
                    connection.execute(text("""
                        INSERT INTO role_descriptions
                            (role_id, version, people_count, role_definition, key_responsibilities, kpis,
                             specifications_experience, education_certifications, minimum_relevant_experience,
                             knowledge_required, skills_required, attitudes_behaviours, ai_refined,
                             validated_by_employee_id, validated_at)
                        VALUES (:role_id, :version, :people_count, :role_definition, :key_responsibilities, :kpis,
                                :specifications_experience, :education_certifications, :minimum_relevant_experience,
                                :knowledge_required, :skills_required, :attitudes_behaviours, FALSE,
                                :validator, :validated_at)
                    """), {**values, "version": version})
                    counts["role_descriptions_created"] += 1
                else:
                    connection.execute(text("""
                        UPDATE role_descriptions SET people_count=:people_count, role_definition=:role_definition,
                            key_responsibilities=:key_responsibilities, kpis=:kpis,
                            specifications_experience=:specifications_experience,
                            education_certifications=:education_certifications,
                            minimum_relevant_experience=:minimum_relevant_experience,
                            knowledge_required=:knowledge_required, skills_required=:skills_required,
                            attitudes_behaviours=:attitudes_behaviours, ai_refined=FALSE,
                            validated_by_employee_id=:validator, validated_at=:validated_at
                        WHERE role_description_id=:role_description_id
                    """), {**values, "role_description_id": existing})
                    counts["role_descriptions_updated"] += 1
                connection.execute(text("""
                    UPDATE roles SET description_status='VALIDATED'
                    WHERE tenant_id=:tenant_id AND role_id=:role_id
                """), {"tenant_id": TENANT_ID, "role_id": role_id})

            for row in source["Employees"]:
                employee_code = clean(row["Employee Code *"])
                employee_id = employees.get(normalized(employee_code))
                if employee_id is None:
                    raise ValueError(f"Employee {employee_code} does not exist; this import only updates supplied employees.")
                division_id = divisions[division_code(row["Division Code *"])]
                function_id = functions[(division_id, normalized(row["Function Name *"]))]
                department_id = departments[(division_id, function_id, normalized(row["Department Name"]))]
                role_id = roles[(division_id, function_id, normalized(row["Role Title"]))]
                manager_code = clean(row.get("Manager Employee Code"))
                manager_id = employees.get(normalized(manager_code)) if manager_code else None
                if manager_code and manager_id is None:
                    raise ValueError(f"Manager {manager_code} is not configured for employee {employee_code}.")
                connection.execute(text("""
                    UPDATE employees SET full_name=:full_name, division_id=:division_id,
                        grade_band_id=:grade_band_id, designation=:designation, function_id=:function_id,
                        department_id=:department_id, role_id=:role_id, manager_employee_id=:manager_employee_id,
                        date_of_joining=:date_of_joining, tenure_start_date=:tenure_start_date,
                        last_promotion_date=:last_promotion_date, retirement_date=:retirement_date,
                        qualification=:qualification, status=:status
                    WHERE tenant_id=:tenant_id AND employee_id=:employee_id
                """), {
                    "tenant_id": TENANT_ID, "employee_id": employee_id,
                    "full_name": clean(row["Full Name *"]), "division_id": division_id,
                    "grade_band_id": grade_bands[normalized(row["Grade Band Code *"])],
                    "designation": clean(row["Designation *"]), "function_id": function_id,
                    "department_id": department_id, "role_id": role_id, "manager_employee_id": manager_id,
                    "date_of_joining": as_date(row.get("Date of Joining")),
                    "tenure_start_date": as_date(row.get("Tenure Start Date")),
                    "last_promotion_date": as_date(row.get("Last Promotion Date")),
                    "retirement_date": as_date(row.get("Retirement Date")),
                    "qualification": clean(row.get("Qualification")) or None,
                    "status": clean(row["Status *"]).upper(),
                })
                counts["employees_updated"] += 1
            if apply:
                transaction.commit()
            else:
                transaction.rollback()
        except Exception:
            transaction.rollback()
            raise
    return counts


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Incrementally import KDDL client master-data workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    load_dotenv(Path(__file__).resolve().parents[1] / ".env")
    print(import_workbook(args.workbook, args.apply))
