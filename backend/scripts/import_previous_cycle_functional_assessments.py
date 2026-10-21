"""Replace closed-cycle Functional Skills assessments from the supplied workbook."""

from __future__ import annotations

import argparse
import os
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from dotenv import load_dotenv
from sqlalchemy import create_engine, text


TENANT_ID = 1
CYCLE_NAME = "FY 2025 Cycle 2"
SUBMITTED_AT = datetime(2026, 3, 31, 17, 0, 0)
DEFAULT_WORKBOOK = Path(r"C:\Users\Bhawna Narula\Downloads\Previous Cycle data-fs.xlsx")
SHEETS = (("D1 — Functional Skills", "D1"), ("D2 — Functional Skills", "D2"))


def clean(value: object) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").split())


def normalized(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(value).casefold())


def rating(value: object) -> int | None:
    if value in (None, ""):
        return None
    try:
        parsed = int(float(value))
    except (TypeError, ValueError) as error:
        raise ValueError(f"Invalid assessment rating {value!r}") from error
    # The workbook uses 0 for an assessment not entered.  Preserve that as a
    # missing rating rather than manufacturing a proficiency level.
    if parsed == 0:
        return None
    if parsed not in {1, 2, 3, 4}:
        raise ValueError(f"Assessment rating must be between 1 and 4, got {value!r}")
    return parsed


def workbook_records(workbook_path: Path) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    records: list[dict[str, object]] = []
    try:
        for sheet_name, division_code in SHEETS:
            sheet = workbook[sheet_name]
            header = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
            skill_columns = [index for index, value in enumerate(header) if clean(value) == "Skill Name"]
            for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
                employee_code = clean(row[1])
                if not employee_code:
                    continue
                for index in skill_columns:
                    skill_name = clean(row[index])
                    if not skill_name:
                        continue
                    description = ""
                    if index + 4 < len(header) and clean(header[index + 4]) == "Description for Improvement":
                        description = clean(row[index + 4])
                    records.append({
                        "source_sheet": sheet_name,
                        "source_row": row_number,
                        "division_code": division_code,
                        "employee_code": employee_code,
                        "employee_name": clean(row[2]),
                        "skill_name": skill_name,
                        "desired_level": rating(row[index + 1]),
                        "self_assessment_level": rating(row[index + 2]),
                        "manager_assessment_level": rating(row[index + 3]),
                        "description_for_improvement": description or None,
                    })
    finally:
        workbook.close()
    return records


def import_assessments(workbook_path: Path, apply: bool) -> dict[str, int]:
    load_dotenv()
    engine = create_engine(os.environ["DATABASE_URL"], pool_pre_ping=True)
    source_records = workbook_records(workbook_path)
    counts: dict[str, int] = defaultdict(int)
    with engine.begin() as connection:
        cycle_id = connection.execute(text("""SELECT cycle_id FROM assessment_cycles
            WHERE tenant_id=:tenant_id AND name=:name AND status='CLOSED'"""),
            {"tenant_id": TENANT_ID, "name": CYCLE_NAME}).scalar_one()
        divisions = {row["code"]: int(row["division_id"]) for row in connection.execute(text(
            "SELECT division_id, code FROM divisions WHERE tenant_id=:tenant_id"), {"tenant_id": TENANT_ID}).mappings()}
        employees = {(int(row["division_id"]), normalized(row["employee_code"])): dict(row) for row in connection.execute(text("""
            SELECT employee_id, division_id, employee_code, role_id, manager_employee_id
            FROM employees WHERE tenant_id=:tenant_id"""), {"tenant_id": TENANT_ID}).mappings()}
        skills = {(int(row["division_id"]), normalized(row["name"])): int(row["skill_id"]) for row in connection.execute(text("""
            SELECT skill_id, division_id, name FROM functional_skills WHERE tenant_id=:tenant_id"""), {"tenant_id": TENANT_ID}).mappings()}

        imported: dict[int, dict[int, dict[str, object]]] = defaultdict(dict)
        for record in source_records:
            counts["source_entries"] += 1
            division_id = divisions[record["division_code"]]
            employee = employees.get((division_id, normalized(record["employee_code"])))
            if employee is None:
                counts["unmatched_employee_entries"] += 1
                continue
            skill_id = skills.get((division_id, normalized(record["skill_name"])))
            if skill_id is None:
                counts["unmatched_skill_entries"] += 1
                continue
            existing = imported[int(employee["employee_id"])].get(skill_id)
            if existing is not None and any(
                existing[field] != record[field]
                for field in ("desired_level", "self_assessment_level", "manager_assessment_level", "description_for_improvement")
            ):
                raise ValueError(f"Conflicting duplicate skill for {record['employee_code']}: {record['skill_name']}")
            imported[int(employee["employee_id"])][skill_id] = {**record, **employee}
            counts["matched_entries"] += 1

        if apply:
            connection.execute(text("""DELETE line_item FROM assessment_line_items line_item
                JOIN assessment_forms form ON form.form_id=line_item.form_id
                WHERE form.tenant_id=:tenant_id AND form.cycle_id=:cycle_id AND form.assessment_type='SKILL'"""),
                {"tenant_id": TENANT_ID, "cycle_id": cycle_id})
            connection.execute(text("""DELETE FROM assessment_forms
                WHERE tenant_id=:tenant_id AND cycle_id=:cycle_id AND assessment_type='SKILL'"""),
                {"tenant_id": TENANT_ID, "cycle_id": cycle_id})

            for employee_id, skill_records in imported.items():
                sample = next(iter(skill_records.values()))
                has_manager_rating = any(item["manager_assessment_level"] is not None for item in skill_records.values())
                form_status = "MANAGER_SUBMITTED" if has_manager_rating else "SELF_SUBMITTED"
                result = connection.execute(text("""INSERT INTO assessment_forms
                    (tenant_id, cycle_id, assessment_type, employee_id, role_id, manager_employee_id, status,
                     self_assessment_date, functional_self_submitted_at, manager_assessment_date, functional_manager_submitted_at)
                    VALUES (:tenant_id,:cycle_id,'SKILL',:employee_id,:role_id,:manager_employee_id,:status,
                     :self_date,:self_submitted_at,:manager_date,:manager_submitted_at)"""), {
                    "tenant_id": TENANT_ID, "cycle_id": cycle_id, "employee_id": employee_id,
                    "role_id": sample["role_id"], "manager_employee_id": sample["manager_employee_id"],
                    "status": form_status, "self_date": SUBMITTED_AT.date(), "self_submitted_at": SUBMITTED_AT,
                    "manager_date": SUBMITTED_AT.date() if has_manager_rating else None,
                    "manager_submitted_at": SUBMITTED_AT if has_manager_rating else None,
                })
                form_id = int(result.lastrowid)
                connection.execute(text("""INSERT INTO assessment_line_items
                    (form_id,item_type,item_id,desired_level,self_assessment_level,manager_assessment_level,description_for_improvement)
                    VALUES (:form_id,'SKILL',:item_id,:desired_level,:self_level,:manager_level,:description)"""), [
                    {"form_id": form_id, "item_id": skill_id, "desired_level": item["desired_level"],
                     "self_level": item["self_assessment_level"], "manager_level": item["manager_assessment_level"],
                     "description": item["description_for_improvement"]}
                    for skill_id, item in skill_records.items()
                ])
            counts["forms_created"] = len(imported)
            counts["line_items_created"] = sum(len(items) for items in imported.values())
        else:
            connection.rollback()
    return dict(counts)


def main() -> None:
    parser = argparse.ArgumentParser(description="Replace previous-cycle Functional Skills assessment records.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--apply", action="store_true")
    arguments = parser.parse_args()
    print(import_assessments(arguments.workbook, arguments.apply))


if __name__ == "__main__":
    main()
