import openpyxl

path = r"C:\Users\Bhawna Narula\VVNT-Founder\OneDrive - VVNT\WIP_Items\Customers - Existing\PGC\FY26\KDDL\KDDL Data\Talent_Client_Data_Template_Populated_FINAL (1).xlsx"
workbook = openpyxl.load_workbook(path, data_only=True)
for name in ("Employees", "Desired Level Matrix", "Functional Skills", "Success Behaviours"):
    sheet = workbook[name]
    rows = [
        [sheet.cell(row, column).value for column in range(1, sheet.max_column + 1)]
        for row in range(2, sheet.max_row + 1)
        if any(sheet.cell(row, column).value not in (None, "") for column in range(1, sheet.max_column + 1))
    ]
    print(f"## {name}: {len(rows)} rows")
    for row in rows[:5]:
        print(row)
